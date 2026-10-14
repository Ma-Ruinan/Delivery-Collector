"""Run in temporary directories only; never touch real evaluation datasets."""
import copy
from contextlib import closing
import json
import sqlite3
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "delivery-collector" / "scripts"))
import collector as c
import xlsx_io as x


def run(status, number=1):
    return {"session_id": f"s{number}", "status": status,
            "submitted_at": f"2026-10-14T09:0{number}:00+08:00",
            "submission_evidence": "message-id/time", "match_evidence": "exact-prompt+uploads",
            "duration": "2m3s", "duration_evidence": "client-duration",
            "reason": "缺少必需文件" if status == "failed" else "完整性检查记录"}


class SelectionTests(unittest.TestCase):
    def test_first_success_priority_even_second_failed_or_running(self):
        for second in ("success", "failed", "running", "unknown"):
            first = run("success")
            result = c.select_run({"attempts": [first, run(second, 2)]}, True)
            self.assertEqual(result[0], "first_success")
            self.assertIs(result[1], first)

    def test_retry_success(self):
        task = {"attempts": [run("failed"), run("success", 2)]}
        decision, selected, _ = c.select_run(task, True)
        self.assertEqual(decision, "second_success")
        value, note = c.excel_values(decision, task, selected)
        self.assertEqual(value, "2m3s")
        self.assertIn("第一次用时2m3s", note)
        self.assertIn("第二次运行成功", note)

    def test_both_failed(self):
        task = {"attempts": [run("failed"), run("failed", 2)]}
        decision, selected, _ = c.select_run(task, True)
        self.assertEqual(decision, "both_failed")
        self.assertEqual(c.excel_values(decision, task, selected)[0], "两次测试均未成功")

    def test_no_runs_only_unexecuted_if_discovery_complete(self):
        self.assertEqual(c.select_run({}, True)[0], "not_executed")
        self.assertEqual(c.select_run({}, False)[0], "pending")

    def test_one_failed_and_more_than_two_pending(self):
        self.assertEqual(c.select_run({"attempts": [run("failed")]}, True)[0], "pending")
        self.assertEqual(c.select_run({"attempts": [run("success", n) for n in (1, 2, 3)]}, True)[0], "pending")

    def test_unknown_first_not_bypassed(self):
        self.assertEqual(c.select_run({"attempts": [run("unknown"), run("success", 2)]}, True)[0], "pending")

    def test_same_session_not_retry(self):
        self.assertEqual(c.select_run({"attempts": [run("failed"), run("success")]}, True)[0], "pending")

    def test_actual_submission_order_checked(self):
        self.assertEqual(c.select_run({"attempts": [run("success", 2), run("success", 1)]}, True)[0], "pending")
        first = run("success"); first["submitted_at"] = "2026-10-14T09:01:00"
        self.assertEqual(c.select_run({"attempts": [first]}, True)[0], "pending")

    def test_missing_timing_not_guessed(self):
        first = run("success"); first.pop("duration_evidence")
        value, note = c.excel_values("first_success", {"attempts": [first]}, first)
        self.assertEqual(value, "")
        self.assertIn("用时无法确认", note)

    def test_normalization_does_not_erase_semantic_differences(self):
        self.assertEqual(c.norm_prompt("题目  \r\n日期2026\n"), "题目\n日期2026")
        self.assertNotEqual(c.norm_prompt("不得使用2025数据"), c.norm_prompt("使用2026数据"))


class FileAndWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self):
        path = self.root / "记录.xlsx"
        x.write_audit(path, {"记录": [["维度", "任务编号", "任务名称", "对象A用时", "对象A说明", "旧对象B用时", "对象B说明"],
                                    ["维度一", "1.1", "问题1", "历史A时间", "历史A说明", "历史B时间", "历史B说明"],
                                    ["", "1.2", "问题2", "", "", "", ""],
                                    ["", "1.3", "问题3", "", "", "", ""]],
                             "其他表": [["保持不变"], ["=不是公式"]]})
        return path

    def plan(self):
        dataset = self.root / "dataset"; dataset.mkdir()
        question = dataset / "维度一" / "问题1"; question.mkdir(parents=True)
        project = self.root / "projects"; project.mkdir()
        source = project / "报告.md"; source.write_text("正式结果", encoding="utf-8")
        workbook = self.workbook(); workbook.rename(dataset / workbook.name); workbook = dataset / workbook.name
        first = run("success")
        first.update(completeness="complete", validation_evidence="实际打开，检查必需项目齐全",
                     artifacts=[{"session_id": "s1", "source": str(source), "sha256": x.digest(source), "evidence": "present_files-success"}])
        return {"schema_version": 1, "dataset_root": str(dataset), "batch_id": "batch1",
                "object": {"name": "WorkBuddy", "version": "5.5.6", "model": "glm-5.2", "prefix": "workbuddy5.5.6"},
                "read_roots": [str(project)], "discovery_complete": True, "discovery_evidence": "全部批次根会话已检查",
                "excel": {"path": str(workbook), "sha256": x.digest(workbook), "sheet": "记录", "task_id_column": "B",
                          "time_column": "F", "note_column": "G", "target_evidence": "用户指定对象B槽位", "overwrite_existing": True,
                          "headers": [{"cell": "F1", "expected": "旧对象B用时", "value": "WorkBuddy-5.5.6 glm-5.2用时"}]},
                "tasks": [{"task_id": "1.1", "question_dir": str(question), "requirements": ["报告"], "attempts": [first, run("success", 2)]}]}

    def test_surgical_update_preserves_all_unrelated_package_parts_and_cells(self):
        path = self.workbook()
        original_bytes = path.read_bytes()
        with zipfile.ZipFile(path) as z:
            before = {n: z.read(n) for n in z.namelist()}
        x.update_workbook(path, [{"sheet": "记录", "cell": "F2", "expected": "历史B时间", "value": "5m6s"}], self.root / "备份.xlsx")
        with zipfile.ZipFile(path) as z:
            for n, payload in before.items():
                if n != "xl/worksheets/sheet1.xml":
                    self.assertEqual(payload, z.read(n))
        cells = x.read_workbook(path)["记录"]["cells"]
        self.assertEqual(cells["D2"]["value"], "历史A时间")
        self.assertEqual(cells["G2"]["value"], "历史B说明")
        self.assertEqual(cells["F2"]["value"], "5m6s")
        self.assertEqual((self.root / "备份.xlsx").read_bytes(), original_bytes)

    def test_stale_cells_rejected_without_modification(self):
        path = self.workbook(); before = x.digest(path)
        with self.assertRaises(ValueError):
            x.update_workbook(path, [{"sheet": "记录", "cell": "F2", "expected": "错误旧值", "value": "new"}], self.root / "备份.xlsx")
        self.assertEqual(x.digest(path), before)
        self.assertFalse((self.root / "备份.xlsx").exists())

    def test_formula_merged_and_protected_rejected(self):
        xml = ('<worksheet xmlns="' + x.NS + '"><sheetData><row r="1"><c r="A1"><f>1+1</f><v>2</v></c></row></sheetData></worksheet>').encode()
        self.assertTrue(x.merged_non_anchor("B2", ["A1:B2"]))
        self.assertFalse(x.merged_non_anchor("A1", ["A1:B2"]))
        path = self.workbook()
        with zipfile.ZipFile(path) as z:
            parts = {i.filename: z.read(i.filename) for i in z.infolist()}
        parts["xl/worksheets/sheet1.xml"] = xml
        with zipfile.ZipFile(path, "w") as z:
            for n, data in parts.items(): z.writestr(n, data)
        with self.assertRaisesRegex(ValueError, "公式"):
            x.update_workbook(path, [{"sheet": "记录", "cell": "A1", "expected": "2", "value": "x"}], self.root / "b.xlsx")
        parts["xl/worksheets/sheet1.xml"] = xml.replace(b"</worksheet>", b'<sheetProtection sheet="1"/></worksheet>')
        with zipfile.ZipFile(path, "w") as z:
            for n, data in parts.items(): z.writestr(n, data)
        with self.assertRaisesRegex(ValueError, "保护"):
            x.update_workbook(path, [{"sheet": "记录", "cell": "A1", "expected": "2", "value": "x"}], self.root / "b.xlsx")

    def test_insert_missing_cell_keeps_order_and_namespace_bytes(self):
        xml = b'<worksheet xmlns="' + x.NS.encode() + b'" xmlns:custom="urn:keep"><sheetData><row r="2"><c r="B2" s="1"><v>1.1000000000000001</v></c><c r="G2"><v>3</v></c></row></sheetData></worksheet>'
        patched = x.patch_sheet(xml, {"F2": "=not-formula"})
        self.assertIn(b'xmlns:custom="urn:keep"', patched)
        self.assertLess(patched.index(b'r="F2"'), patched.index(b'r="G2"'))
        self.assertIn(b'1.1000000000000001', patched)

    def test_execute_single_object_and_idempotence(self):
        plan = self.plan(); first = c.execute(plan, True)
        self.assertEqual(first["copied"], 1)
        self.assertEqual(first["excel"]["status"], "updated")
        cells = x.read_workbook(plan["excel"]["path"])["记录"]["cells"]
        self.assertEqual(cells["D2"]["value"], "历史A时间")
        self.assertEqual(cells["F2"]["value"], "2m3s")
        self.assertEqual(cells["G2"]["value"], "一次运行成功")
        self.assertEqual(cells["F1"]["value"], "WorkBuddy-5.5.6 glm-5.2用时")
        plan["excel"]["sha256"] = x.digest(plan["excel"]["path"])
        plan["excel"]["headers"][0]["expected"] = cells["F1"]["value"]
        second = c.execute(plan, True)
        self.assertEqual(second["already_present"], 1)
        self.assertEqual(second["copied"], 0)
        self.assertEqual(second["excel"]["status"], "unchanged")

    def test_first_success_not_blocked_by_extra_failure_missing_reason(self):
        plan = self.plan()
        extra = run("failed", 2); extra.pop("reason")
        plan["tasks"][0]["attempts"][1] = extra
        result = c.execute(plan, True)
        self.assertEqual(result["copied"], 1)
        self.assertEqual(result["excel"]["status"], "updated")

    def test_preflight_does_not_copy_or_update(self):
        plan = self.plan(); before = x.digest(plan["excel"]["path"])
        result = c.execute(plan, False)
        self.assertEqual(result["copied"], 0)
        self.assertEqual(x.digest(plan["excel"]["path"]), before)
        self.assertFalse(any(Path(plan["tasks"][0]["question_dir"]).iterdir()))
        self.assertTrue(Path(result["checklist"]).is_file())

    def test_pending_task_does_not_block_other_task(self):
        plan = self.plan(); bad = copy.deepcopy(plan["tasks"][0]); bad["task_id"] = "1.2"; bad["pending"] = "来源不明确"
        plan["tasks"].insert(0, bad)
        result = c.execute(plan, True)
        self.assertEqual(result["copied"], 1)
        self.assertTrue(any(i["task_id"] == "1.2" for i in result["issues"]))
        self.assertEqual(x.read_workbook(plan["excel"]["path"])["记录"]["cells"]["F3"]["value"], "")

    def test_unexecuted_preserves_historical_results_and_empty_delivery(self):
        plan = self.plan(); plan["tasks"][0]["attempts"] = []
        before = x.digest(plan["excel"]["path"])
        result = c.execute(plan, True)
        self.assertEqual(result["copied"], 0)
        self.assertEqual(x.digest(plan["excel"]["path"]), before)
        audit = json.loads(Path(result["audit_json"]).read_text(encoding="utf-8"))
        self.assertEqual(audit["tasks"][0]["decision"], "not_executed")

    def test_target_conflict_not_overwritten_or_recorded_success(self):
        plan = self.plan(); target = Path(plan["tasks"][0]["question_dir"]) / "workbuddy5.5.6-1.1-报告.md"
        target.write_text("历史不同结果", encoding="utf-8")
        result = c.execute(plan, True)
        self.assertEqual(result["copied"], 0)
        self.assertEqual(target.read_text(encoding="utf-8"), "历史不同结果")
        self.assertEqual(result["excel"]["status"], "not_updated")

    def test_changed_source_hash_not_collected(self):
        plan = self.plan(); Path(plan["tasks"][0]["attempts"][0]["artifacts"][0]["source"]).write_text("changed", encoding="utf-8")
        self.assertEqual(c.execute(plan, True)["copied"], 0)

    def test_outside_question_directory_rejected(self):
        plan = self.plan(); plan["tasks"][0]["question_dir"] = str(self.root)
        self.assertEqual(c.execute(plan, True)["copied"], 0)

    def test_project_dependency_relative_paths_preserved(self):
        plan = self.plan(); first = plan["tasks"][0]["attempts"][0]
        project = Path(plan["read_roots"][0]); (project / "css").mkdir()
        (project / "index.html").write_text('<link href="css/style.css">内容', encoding="utf-8")
        (project / "css" / "style.css").write_text("body {color:black}", encoding="utf-8")
        first["artifacts"] = [{"session_id": "s1", "source": str(project / rel), "sha256": x.digest(project / rel),
                               "evidence": "正式项目依赖", "project_name": "site", "relative_path": rel} for rel in ("index.html", "css/style.css")]
        result = c.execute(plan, True)
        self.assertEqual(result["copied"], 2)
        base = Path(plan["tasks"][0]["question_dir"]) / "workbuddy5.5.6-1.1-site"
        self.assertEqual((base / "index.html").read_bytes(), (project / "index.html").read_bytes())
        self.assertTrue((base / "css" / "style.css").exists())

    def test_second_run_docx_only_excludes_thinking_preserves_formal_text(self):
        from docx import Document
        plan = self.plan(); task = plan["tasks"][0]
        task["attempts"] = [run("failed"), run("failed", 2)]
        second = task["attempts"][1]
        second.update(official_output_complete=True, official_output=[
            {"kind": "thinking", "text": "隐藏推理不得进入DOCX"},
            {"kind": "text", "visible": True, "text": "正式回答\n思考这个业务问题仍属于正式文本\n可见错误", "evidence": "part-formal"},
            {"kind": "table", "visible": True, "rows": [["字段", "值"], ["结果", "部分"]], "evidence": "formal-table"}])
        result = c.execute(plan, True)
        self.assertEqual(result["copied"], 1)
        files = list(Path(task["question_dir"]).iterdir())
        self.assertEqual([f.name for f in files], ["workbuddy5.5.6-1.1-未成功任务界面输出.docx"])
        doc = Document(files[0]); text = "\n".join(p.text for p in doc.paragraphs)
        self.assertEqual(text, "正式回答\n思考这个业务问题仍属于正式文本\n可见错误")
        self.assertEqual(doc.tables[0].cell(1, 1).text, "部分")
        self.assertEqual(x.read_workbook(plan["excel"]["path"])["记录"]["cells"]["F2"]["value"], "两次测试均未成功")
        plan["excel"]["sha256"] = x.digest(plan["excel"]["path"])
        plan["excel"]["headers"][0]["expected"] = "WorkBuddy-5.5.6 glm-5.2用时"
        self.assertEqual(c.execute(plan, True)["already_present"], 1)

    def test_missing_formal_output_not_invented(self):
        plan = self.plan(); plan["tasks"][0]["attempts"] = [run("failed"), run("failed", 2)]
        result = c.execute(plan, True)
        self.assertEqual(result["copied"], 0)
        self.assertTrue(any("正式界面输出" in i["issue"] for i in result["issues"]))

    def test_stale_workbook_does_not_prevent_collecting_files(self):
        plan = self.plan(); plan["excel"]["sha256"] = "wrong"
        result = c.execute(plan, True)
        self.assertEqual(result["copied"], 1)
        self.assertEqual(result["excel"]["status"], "not_updated")
        self.assertTrue(result["issues"])

    def test_roster_and_readonly_sqlite(self):
        question = self.root / "维度二" / "问题3"; question.mkdir(parents=True)
        (question / "问题描述.txt").write_text("题目", encoding="utf-8")
        self.assertEqual(len(c.roster(str(self.root))["tasks"]), 1)
        db = self.root / "app.db"
        with closing(sqlite3.connect(db)) as con:
            con.execute("CREATE TABLE session (id TEXT, title TEXT)")
            con.execute("INSERT INTO session VALUES ('s1', '不依赖标题')")
            con.commit()
        before = x.digest(db)
        result = c.sqlite_read(SimpleNamespace(db=str(db), table="session", equals=["id", "s1"], limit=1, offset=0))
        self.assertEqual(result["rows"][0]["id"], "s1")
        self.assertEqual(x.digest(db), before)

    def test_excel_strings_cannot_inject_formulas(self):
        path = self.workbook()
        cells = x.read_workbook(path)["其他表"]["cells"]
        self.assertEqual(cells["A2"]["value"], "=不是公式")
        self.assertIsNone(cells["A2"]["formula"])


if __name__ == "__main__":
    unittest.main()
